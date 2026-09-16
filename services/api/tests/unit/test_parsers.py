"""Parser tests build real files rather than mocking the libraries.

A parser's job is to survive whatever a client uploads, so the fixtures here are produced
the same way a client's file would be: written by a document library, not hand-crafted.
"""

import io

import pytest

from harbor_api.ingestion.parsers import (
    MAX_UPLOAD_BYTES,
    UnparsableDocument,
    parse,
    validate_upload,
)


def make_pdf(pages: list[str]) -> bytes:
    """A PDF with a real text layer. pypdf cannot author text, so reportlab draws it."""
    pytest.importorskip("reportlab", reason="reportlab builds the text PDF fixture")
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for text in pages:
        pdf.drawString(72, 720, text)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    import docx

    document = docx.Document()
    for style, text in blocks:
        document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestValidation:
    def test_rejects_empty_file(self) -> None:
        with pytest.raises(UnparsableDocument, match="empty"):
            validate_upload(b"", "policy.pdf")

    def test_rejects_oversized_file(self) -> None:
        with pytest.raises(UnparsableDocument, match="larger than"):
            validate_upload(b"%PDF-" + b"x" * MAX_UPLOAD_BYTES, "policy.pdf")

    def test_rejects_unsupported_extension(self) -> None:
        with pytest.raises(UnparsableDocument, match="Unsupported file type"):
            validate_upload(b"MZ\x90\x00", "installer.exe")

    def test_rejects_a_file_whose_bytes_contradict_its_extension(self) -> None:
        # Renaming handbook.exe to handbook.pdf must not get it parsed.
        with pytest.raises(UnparsableDocument, match="does not look like a real PDF"):
            validate_upload(b"MZ\x90\x00 not a pdf at all", "handbook.pdf")

    def test_accepts_a_real_pdf(self) -> None:
        assert validate_upload(b"%PDF-1.7\n...", "handbook.pdf") == ".pdf"


class TestMarkdownAndText:
    def test_markdown_is_passed_through(self) -> None:
        parsed = parse("# 員工手冊\n\n年假七天。".encode(), "handbook.md")

        assert parsed.markdown.startswith("# 員工手冊")

    def test_big5_encoded_text_is_decoded(self) -> None:
        # HK clients still export Big5 from older systems.
        parsed = parse("僱員手冊".encode("big5"), "handbook.txt")

        assert parsed.markdown == "僱員手冊"


class TestHtml:
    def test_headings_lists_and_paragraphs_become_markdown(self) -> None:
        html = b"""
        <html><body>
          <nav>skip me</nav>
          <h1>Customer FAQ</h1>
          <p>We open at 07:00.</p>
          <h2>Delivery</h2>
          <ul><li>Free above HK$500</li></ul>
          <script>ignored()</script>
        </body></html>
        """

        markdown = parse(html, "faq.html").markdown

        assert "# Customer FAQ" in markdown
        assert "## Delivery" in markdown
        assert "- Free above HK$500" in markdown
        assert "skip me" not in markdown
        assert "ignored" not in markdown

    def test_html_without_text_is_rejected(self) -> None:
        with pytest.raises(UnparsableDocument, match="No readable text"):
            parse(b"<html><body><script>x()</script></body></html>", "empty.html")


class TestPdf:
    def test_text_pdf_is_extracted_with_page_count(self) -> None:
        parsed = parse(make_pdf(["Probation lasts three months.", "Page two."]), "handbook.pdf")

        assert "Probation lasts three months." in parsed.markdown
        assert parsed.page_count == 2

    def test_scanned_pdf_fails_with_an_actionable_message(self) -> None:
        # A blank page stands in for a scan: pypdf extracts nothing from either.
        with pytest.raises(UnparsableDocument, match="scanned PDF"):
            parse(make_blank_pdf(), "scan.pdf")

    def test_corrupt_pdf_is_rejected(self) -> None:
        with pytest.raises(UnparsableDocument):
            parse(b"%PDF-1.7\nbroken", "broken.pdf")


class TestDocx:
    def test_headings_and_body_become_markdown(self) -> None:
        data = make_docx(
            [
                ("Heading 1", "員工手冊"),
                ("Heading 2", "病假"),
                ("", "請病假需要醫生證明書。"),
            ]
        )

        markdown = parse(data, "handbook.docx").markdown

        assert "# 員工手冊" in markdown
        assert "## 病假" in markdown
        assert "請病假需要醫生證明書。" in markdown

    def test_empty_document_is_rejected(self) -> None:
        with pytest.raises(UnparsableDocument, match="No readable text"):
            parse(make_docx([]), "empty.docx")


class TestEncodingDetection:
    """Legacy Chinese encodings decode each other's bytes without raising, so the decoder
    scores candidates instead of trusting the first one that does not throw."""

    @pytest.mark.parametrize(
        ("text", "encoding"),
        [
            ("僱員手冊：病假需要醫生證明書。", "big5"),
            ("员工手册：病假需要医生证明书。", "gb18030"),
            ("Employee handbook: sick leave", "utf-8"),
            ("員工手冊 with English", "utf-8"),
        ],
    )
    def test_round_trips_common_encodings(self, text: str, encoding: str) -> None:
        assert parse(text.encode(encoding), "handbook.txt").markdown == text

    def test_utf16_with_a_byte_order_mark_is_honoured(self) -> None:
        assert parse("年假七天".encode("utf-16"), "handbook.txt").markdown == "年假七天"
